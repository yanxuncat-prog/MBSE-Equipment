"""Report reverse parsing: Word/PDF upload, field extraction, data backfill."""
import re
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.user import User
from app.models.equipment import Equipment
from app.models.configuration import ConfigEquipment
from app.api.deps import get_current_user

router = APIRouter(prefix="/report-parsing", tags=["report-parsing"])


# ── Pydantic schemas ────────────────────────────────────────────────

class FieldApply(BaseModel):
    field_name: str
    value: str


class ApplyRequest(BaseModel):
    config_id: str
    equipment_id: str
    fields: list[FieldApply]


# ── Text extraction ─────────────────────────────────────────────────

def _extract_text_from_docx(content: bytes) -> str:
    import docx
    from io import BytesIO
    doc = docx.Document(BytesIO(content))
    parts: list[str] = []
    # Paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    # Tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_text_from_pdf(content: bytes) -> str:
    import pdfplumber
    from io import BytesIO
    parts: list[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)


# ── Field extraction (regex) ────────────────────────────────────────

EXTRACTION_RULES: list[dict[str, Any]] = [
    {
        "field_name": "part_number",
        "display": "件号",
        "patterns": [
            r"件号[：:]\s*(.+?)(?:\s*$|\s{2,})",
            r"Part\s*Number[：:]\s*(.+?)(?:\s*$|\s{2,})",
            r"P/N[：:]\s*(.+?)(?:\s*$|\s{2,})",
        ],
    },
    {
        "field_name": "name",
        "display": "设备名称",
        "patterns": [
            r"设备名称[：:]\s*(.+?)(?:\s*$|\s{2,})",
            r"名称[：:]\s*(.+?)(?:\s*$|\s{2,})",
            r"Equipment\s*Name[：:]\s*(.+?)(?:\s*$|\s{2,})",
        ],
    },
    {
        "field_name": "mass_kg",
        "display": "重量",
        "patterns": [
            r"重量[：:]\s*([\d.]+)\s*kg",
            r"Weight[：:]\s*([\d.]+)\s*kg",
            r"质量[：:]\s*([\d.]+)\s*kg",
        ],
    },
    {
        "field_name": "ata_chapter",
        "display": "ATA章节",
        "patterns": [
            r"ATA[：:]\s*(\d+[-\s]?\d*)",
            r"ATA\s*Chapter[：:]\s*(\d+[-\s]?\d*)",
            r"ATA章节[：:]\s*(\d+[-\s]?\d*)",
        ],
    },
    {
        "field_name": "dimensions_mm",
        "display": "尺寸",
        "patterns": [
            r"尺寸[：:]\s*(.+?)(?:mm|$)",
            r"Dimensions[：:]\s*(.+?)(?:mm|$)",
            r"外形尺寸[：:]\s*(.+?)(?:mm|$)",
        ],
    },
    {
        "field_name": "voltage_range",
        "display": "工作电压",
        "patterns": [
            r"工作电压[：:]\s*(.+?)(?:V|$)",
            r"电压范围[：:]\s*(.+?)(?:V|$)",
            r"Voltage[：:]\s*(.+?)(?:V|$)",
        ],
    },
    {
        "field_name": "power_watts",
        "display": "用电功率",
        "patterns": [
            r"功率[：:]\s*(.+?)(?:W|$)",
            r"用电功率[：:]\s*(.+?)(?:W|$)",
            r"Power[：:]\s*(.+?)(?:W|$)",
        ],
    },
    {
        "field_name": "dal",
        "display": "DAL等级",
        "patterns": [
            r"DAL[：:]\s*([A-E])",
            r"设计保证等级[：:]\s*([A-E])",
        ],
    },
]


def _extract_fields(text: str) -> list[dict]:
    results = []
    for rule in EXTRACTION_RULES:
        matched = False
        for pattern in rule["patterns"]:
            m = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
            if m:
                value = m.group(1).strip()
                if value:
                    # Determine confidence based on pattern specificity
                    confidence = "high" if rule["patterns"].index(pattern) == 0 else "medium"
                    results.append({
                        "field_name": rule["field_name"],
                        "field_value": value,
                        "confidence": confidence,
                    })
                    matched = True
                    break
        # If no exact pattern matched, try a looser match
        if not matched:
            loose = re.search(
                rf'{rule["display"]}[\s：:]+(.{{2,50}})',
                text,
                re.IGNORECASE,
            )
            if loose:
                results.append({
                    "field_name": rule["field_name"],
                    "field_value": loose.group(1).strip(),
                    "confidence": "low",
                })
    return results


# ── Field mapping for apply ──────────────────────────────────────────

EQUIPMENT_FIELDS = {
    "part_number", "name", "ata_chapter", "equipment_type", "description",
    "name_en", "abbreviation_en",
    "supplier_part_number", "dal", "dimensions_mm",
    "voltage_range", "power_redundancy", "power_voltage", "power_watts",
    "shell_grounding_method",
}
CONFIG_EQUIPMENT_FIELDS = {
    "mass_kg", "power_kva_normal", "power_kva_emergency", "power_kva_max",
    "install_method", "bonding_method", "bonding_type", "bonding_resistance",
}


# ── Endpoints ────────────────────────────────────────────────────────

@router.post("/upload")
async def upload_and_parse(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
):
    if not file.filename:
        raise HTTPException(400, "未提供文件名")

    ext = Path(file.filename).suffix.lower()
    if ext not in (".docx", ".pdf"):
        raise HTTPException(400, f"不支持的文件类型: {ext}. 仅支持 .docx 和 .pdf")

    content = await file.read()
    if len(content) > 50 * 1024 * 1024:  # 50MB limit
        raise HTTPException(400, "文件过大，最大 50MB")

    try:
        if ext == ".docx":
            text = _extract_text_from_docx(content)
        else:
            text = _extract_text_from_pdf(content)
    except Exception as e:
        raise HTTPException(400, f"文件解析失败: {str(e)}")

    parsed_fields = _extract_fields(text)
    raw_preview = text[:500] if len(text) > 500 else text

    return {
        "parsed_fields": parsed_fields,
        "raw_text_preview": raw_preview,
    }


@router.post("/apply")
async def apply_parsed_fields(
    body: ApplyRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    applied = 0
    skipped = 0

    # Load equipment
    eq_result = await db.execute(select(Equipment).where(Equipment.id == body.equipment_id))
    equipment = eq_result.scalar_one_or_none()
    if not equipment:
        raise HTTPException(404, "设备不存在")

    # Load config_equipment
    ce_result = await db.execute(
        select(ConfigEquipment).where(
            ConfigEquipment.config_id == body.config_id,
            ConfigEquipment.equipment_id == body.equipment_id,
        )
    )
    config_equip = ce_result.scalar_one_or_none()

    for field in body.fields:
        fname = field.field_name
        value = field.value

        if fname in EQUIPMENT_FIELDS:
            if hasattr(equipment, fname):
                setattr(equipment, fname, value)
                applied += 1
            else:
                skipped += 1
        elif fname in CONFIG_EQUIPMENT_FIELDS:
            if config_equip and hasattr(config_equip, fname):
                # Convert numeric fields
                try:
                    col_type = type(getattr(ConfigEquipment, fname).property.columns[0].type)
                    if "Float" in str(col_type) or "float" in str(col_type):
                        value = float(value)
                except (ValueError, AttributeError):
                    pass
                setattr(config_equip, fname, value)
                applied += 1
            else:
                skipped += 1
        else:
            skipped += 1

    await db.commit()
    return {"applied_count": applied, "skipped_count": skipped}
