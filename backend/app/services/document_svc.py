import io
import uuid
from datetime import datetime, timezone
from collections import defaultdict

from jinja2 import Environment, FileSystemLoader
from openpyxl import Workbook

from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# WeasyPrint requires system libraries (pango/glib) -- lazy import to avoid crash if missing
def _get_weasyprint_html():
    from weasyprint import HTML
    return HTML
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Equipment, Configuration, BusDefinition, Zone, ConfigEquipment
from app.models.configuration import ConfigEquipment as ConfigEquipmentModel
from app.services.constraint_svc import validate_config


TEMPLATE_DIR = "templates"
jinja_env = Environment(loader=FileSystemLoader(TEMPLATE_DIR))


async def _load_config_equipment(db: AsyncSession, config_id: str) -> tuple[Configuration, list[dict]]:
    """Load config and its equipment with config-level position/bus data."""
    config = await db.get(Configuration, config_id)
    result = await db.execute(
        select(ConfigEquipmentModel)
        .options(
            selectinload(ConfigEquipmentModel.equipment).selectinload(Equipment.supplier),
            selectinload(ConfigEquipmentModel.zone),
            selectinload(ConfigEquipmentModel.bus),
        )
        .where(ConfigEquipmentModel.config_id == config_id)
    )
    ce_list = list(result.scalars().unique().all())

    # Build enriched dicts for each equipment
    items = []
    for ce in ce_list:
        e = ce.equipment
        items.append({
            "equipment": e,
            "zone_id": ce.zone_id,
            "zone_code": ce.zone.zone_code if ce.zone else None,
            "zone_name": ce.zone.name if ce.zone else None,
            "sta": ce.sta,
            "wl": ce.wl,
            "bl": ce.bl,
            "rack_position": ce.rack_position,
            "bus_id": ce.bus_id,
            "bus_name": ce.bus.bus_name if ce.bus else None,
            "install_method": ce.install_method,
            "layout_adjustment": ce.layout_adjustment,
            "mass_kg": ce.mass_kg,
            "cg_x": ce.cg_x,
            "cg_y": ce.cg_y,
            "cg_z": ce.cg_z,
        })

    return config, items


async def generate_equipment_list_pdf(db: AsyncSession, config_id: str) -> bytes:
    config, items = await _load_config_equipment(db, config_id)

    grouped = defaultdict(list)
    total_mass = 0.0
    for item in items:
        e = item["equipment"]
        mass = item.get("mass_kg")
        row = {
            "part_number": e.part_number, "name": e.name, "equipment_type": e.equipment_type,
            "status": e.status,
            "mass_kg": f"{mass:.1f}" if mass else None,
            "zone_code": item["zone_code"],
            "sta": f"{item['sta']:.0f}" if item["sta"] is not None else None,
            "bus_name": item["bus_name"],
            "power_kva": f"{e.power_kva_normal:.2f}" if e.power_kva_normal else None,
        }
        ata = e.ata_chapter.split("-")[0] if "-" in e.ata_chapter else e.ata_chapter
        grouped[ata].append(row)
        if mass:
            total_mass += mass

    template = jinja_env.get_template("equipment_list.html")
    html_str = template.render(
        config_version=config.version,
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        grouped_by_ata=dict(sorted(grouped.items())),
        total_count=len(items),
        total_mass_kg=f"{total_mass:.1f}",
    )
    return _get_weasyprint_html()(string=html_str).write_pdf()


async def generate_equipment_list_xlsx(db: AsyncSession, config_id: str) -> bytes:
    config, items = await _load_config_equipment(db, config_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "设备清单"
    ws.append(["件号", "名称", "ATA章节", "类型", "重量(kg)", "区域", "STA", "母线", "功耗(kVA)", "状态"])

    sorted_items = sorted(items, key=lambda x: (x["equipment"].ata_chapter, x["equipment"].part_number))
    for item in sorted_items:
        e = item["equipment"]
        ws.append([
            e.part_number, e.name, e.ata_chapter, e.equipment_type,
            item.get("mass_kg"),
            item["zone_code"],
            item["sta"],
            item["bus_name"],
            e.power_kva_normal,
            e.status,
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def generate_weight_report_pdf(db: AsyncSession, config_id: str) -> bytes:
    config, items = await _load_config_equipment(db, config_id)
    report = await validate_config(db, config_id)

    wb_engine_result = next((e for e in report.engines if e.engine_name == "weight_balance"), None)
    details = wb_engine_result.details if wb_engine_result else {}

    equipment_rows = []
    sorted_items = sorted(items, key=lambda x: x["equipment"].ata_chapter)
    for item in sorted_items:
        e = item["equipment"]
        mass = item.get("mass_kg")
        if mass:
            arm_sta = item["sta"] if item["sta"] is not None else 0.0
            equipment_rows.append({
                "part_number": e.part_number, "name": e.name, "ata_chapter": e.ata_chapter,
                "mass_kg": f"{mass:.1f}",
                "arm_sta": f"{arm_sta:.0f}",
                "moment": f"{mass * arm_sta:.0f}",
            })

    status = wb_engine_result.status if wb_engine_result else "pass"
    template = jinja_env.get_template("weight_report.html")
    html_str = template.render(
        config_version=config.version,
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        total_mass_kg=f"{details.get('total_mass_kg', 0):.1f}",
        cg_sta=f"{details.get('cg_sta', 0):.1f}",
        cg_pct_mac=f"{details.get('cg_pct_mac', 0):.1f}",
        mtow_margin_kg=f"{details.get('mtow_margin_kg', 0):.1f}",
        mtow_ratio_pct=f"{details.get('mtow_ratio_pct', 0):.1f}",
        status_class=status,
        status_text={"pass": "✓ 包线内", "warning": "⚠ 接近限值", "blocked": "✗ 超限"}.get(status, status),
        equipment=equipment_rows,
    )
    return _get_weasyprint_html()(string=html_str).write_pdf()


async def generate_eload_report_pdf(db: AsyncSession, config_id: str, phase: str = "normal") -> bytes:
    config, items = await _load_config_equipment(db, config_id)
    report = await validate_config(db, config_id, phase=phase)

    eload_result = next((e for e in report.engines if e.engine_name == "electrical_load"), None)
    bus_details = eload_result.details.get("buses", {}) if eload_result else {}

    bus_result = await db.execute(select(BusDefinition).where(BusDefinition.program_id == config.program_id))
    bus_defs = {str(b.id): b for b in bus_result.scalars().all()}

    # Build a lookup: bus_id -> list of equipment items assigned to that bus in this config
    bus_equip_map = defaultdict(list)
    for item in items:
        if item["bus_id"] and item["equipment"].power_kva_normal is not None:
            bus_equip_map[str(item["bus_id"])].append(item)

    buses = []
    for bid, info in bus_details.items():
        bdef = bus_defs.get(bid)
        ratio = info.get("load_ratio_pct", 0)
        status_class = "blocked" if ratio > 100 else "warning" if ratio > 85 else "pass"
        status_text = "过载" if ratio > 100 else "接近满载" if ratio > 85 else "正常"

        bus_equip = []
        for item in bus_equip_map.get(bid, []):
            e = item["equipment"]
            bus_equip.append({
                "part_number": e.part_number, "name": e.name, "ata_chapter": e.ata_chapter,
                "power_normal": f"{e.power_kva_normal:.2f}" if e.power_kva_normal else "0.00",
                "power_emergency": f"{e.power_kva_emergency:.2f}" if e.power_kva_emergency else None,
                "power_max": f"{e.power_kva_max:.2f}" if e.power_kva_max else None,
            })

        buses.append({
            "bus_name": info.get("bus_name", ""), "bus_type": bdef.bus_type if bdef else "",
            "capacity_kva": f"{info.get('capacity_kva', 0):.1f}",
            "load_kva": f"{info.get('load_kva', 0):.1f}",
            "load_ratio_pct": f"{ratio:.1f}",
            "margin_kva": f"{info.get('margin_kva', 0):.1f}",
            "status_class": status_class, "status_text": status_text,
            "equipment": bus_equip,
        })

    template = jinja_env.get_template("eload_report.html")
    html_str = template.render(
        config_version=config.version, phase=phase,
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        buses=buses,
    )
    return _get_weasyprint_html()(string=html_str).write_pdf()


# ---------------------------------------------------------------------------
# Word (docx) helpers
# ---------------------------------------------------------------------------

def _set_cell_text(cell, text: str, bold: bool = False, size: int = 10, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Set cell text with formatting."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text) if text is not None else "-")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _style_table(table):
    """Apply a clean bordered style to a docx table."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single",
            qn("w:sz"): "4",
            qn("w:space"): "0",
            qn("w:color"): "000000",
        })
        borders.append(el)
    tbl_pr.append(borders)


def _add_info_row(table, label: str, value):
    """Add a two-column label/value row to a table."""
    row = table.add_row()
    _set_cell_text(row.cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    _set_cell_text(row.cells[1], str(value) if value is not None else "-", align=WD_ALIGN_PARAGRAPH.LEFT)


async def generate_installation_report_docx(db: AsyncSession, config_id: str) -> bytes:
    """Generate a Word installation report for a configuration."""
    config, items = await _load_config_equipment(db, config_id)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    doc = DocxDocument()

    # -- Page margins --
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ===== Title page =====
    for _ in range(6):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("设备安装报告")
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_paragraph("")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"构型版本: {config.version}")
    run.font.size = Pt(16)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_paragraph("")

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info_p.add_run(f"生成日期: {generated_at}")
    run.font.size = Pt(12)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    count_p = doc.add_paragraph()
    count_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = count_p.add_run(f"设备总数: {len(items)}")
    run.font.size = Pt(12)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_page_break()

    # ===== Equipment summary table =====
    doc.add_heading("设备汇总表", level=1)

    summary_headers = ["序号", "设备名称", "件号", "ATA", "重量(kg)", "STA位置", "安装方式"]
    summary_table = doc.add_table(rows=1, cols=len(summary_headers))
    _style_table(summary_table)

    # Header row
    for i, header in enumerate(summary_headers):
        _set_cell_text(summary_table.rows[0].cells[i], header, bold=True, size=9)

    sorted_items = sorted(items, key=lambda x: (x["equipment"].ata_chapter, x["equipment"].part_number))
    for idx, item in enumerate(sorted_items, 1):
        e = item["equipment"]
        mass = item.get("mass_kg")

        row = summary_table.add_row()
        _set_cell_text(row.cells[0], str(idx), size=9)
        _set_cell_text(row.cells[1], e.name, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
        _set_cell_text(row.cells[2], e.part_number, size=9)
        _set_cell_text(row.cells[3], e.ata_chapter, size=9)
        _set_cell_text(row.cells[4], f"{mass:.1f}" if mass else "-", size=9)
        _set_cell_text(row.cells[5], f"{item['sta']:.0f}" if item["sta"] is not None else "-", size=9)
        _set_cell_text(row.cells[6], item.get("install_method") or "-", size=9, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()

    # ===== Per-equipment detail sections =====
    doc.add_heading("设备详细信息", level=1)

    for idx, item in enumerate(sorted_items, 1):
        e = item["equipment"]

        doc.add_heading(f"{idx}. {e.name}", level=2)

        # --- Basic info table ---
        doc.add_heading("基本信息", level=3)
        basic_table = doc.add_table(rows=0, cols=2)
        _style_table(basic_table)
        basic_table.columns[0].width = Cm(5)
        basic_table.columns[1].width = Cm(11)

        _add_info_row(basic_table, "件号", e.part_number)
        _add_info_row(basic_table, "ATA", e.ata_chapter)
        _add_info_row(basic_table, "类型", e.equipment_type)
        _add_info_row(basic_table, "状态", e.status)
        _add_info_row(basic_table, "供电电压", e.power_voltage)

        doc.add_paragraph("")

        # --- Weight / CG info ---
        doc.add_heading("重量/重心", level=3)
        wt_table = doc.add_table(rows=0, cols=2)
        _style_table(wt_table)
        wt_table.columns[0].width = Cm(5)
        wt_table.columns[1].width = Cm(11)

        mass_val = item.get("mass_kg")
        _add_info_row(wt_table, "重量(kg)", f"{mass_val:.1f}" if mass_val else "-")
        cg_x = item.get("cg_x")
        cg_y = item.get("cg_y")
        cg_z = item.get("cg_z")
        _add_info_row(wt_table, "重心X(mm)", f"{cg_x:.1f}" if cg_x is not None else "-")
        _add_info_row(wt_table, "重心Y(mm)", f"{cg_y:.1f}" if cg_y is not None else "-")
        _add_info_row(wt_table, "重心Z(mm)", f"{cg_z:.1f}" if cg_z is not None else "-")

        doc.add_paragraph("")

        # --- Installation info ---
        doc.add_heading("安装信息", level=3)
        inst_table = doc.add_table(rows=0, cols=2)
        _style_table(inst_table)
        inst_table.columns[0].width = Cm(5)
        inst_table.columns[1].width = Cm(11)

        _add_info_row(inst_table, "STA", f"{item['sta']:.0f}" if item["sta"] is not None else "-")
        _add_info_row(inst_table, "WL", f"{item['wl']:.0f}" if item["wl"] is not None else "-")
        _add_info_row(inst_table, "BL", f"{item['bl']:.0f}" if item["bl"] is not None else "-")
        _add_info_row(inst_table, "安装方式", item.get("install_method") or "-")
        _add_info_row(inst_table, "布置调整", item.get("layout_adjustment") or "-")

        # Add a page break between equipment except for the last one
        if idx < len(sorted_items):
            doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def generate_equipment_list_docx(db: AsyncSession, config_id: str) -> bytes:
    """Generate a Word equipment list document (same content as xlsx but in docx)."""
    config, items = await _load_config_equipment(db, config_id)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    doc = DocxDocument()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    doc.add_heading(f"设备清单 — {config.version}", level=1)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = meta.add_run(f"生成日期: {generated_at}    设备总数: {len(items)}")
    run.font.size = Pt(10)

    headers = ["序号", "件号", "名称", "ATA", "类型", "重量(kg)", "区域", "STA", "母线", "功耗(kVA)", "状态"]
    table = doc.add_table(rows=1, cols=len(headers))
    _style_table(table)

    for i, h in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], h, bold=True, size=8)

    sorted_items = sorted(items, key=lambda x: (x["equipment"].ata_chapter, x["equipment"].part_number))
    for idx, item in enumerate(sorted_items, 1):
        e = item["equipment"]
        row = table.add_row()
        _set_cell_text(row.cells[0], str(idx), size=8)
        _set_cell_text(row.cells[1], e.part_number, size=8)
        _set_cell_text(row.cells[2], e.name, size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
        _set_cell_text(row.cells[3], e.ata_chapter, size=8)
        _set_cell_text(row.cells[4], e.equipment_type, size=8)
        _set_cell_text(row.cells[5], f"{item.get('mass_kg'):.1f}" if item.get("mass_kg") else "-", size=8)
        _set_cell_text(row.cells[6], item["zone_code"] or "-", size=8)
        _set_cell_text(row.cells[7], f"{item['sta']:.0f}" if item["sta"] is not None else "-", size=8)
        _set_cell_text(row.cells[8], item["bus_name"] or "-", size=8)
        _set_cell_text(row.cells[9], f"{e.power_kva_normal:.2f}" if e.power_kva_normal else "-", size=8)
        _set_cell_text(row.cells[10], e.status, size=8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
