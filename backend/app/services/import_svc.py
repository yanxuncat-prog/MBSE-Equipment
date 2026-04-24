import uuid
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Equipment, ElectricalLoad


COLUMN_MAP = {
    "件号": "part_number",
    "名称": "name",
    "ATA章节": "ata_chapter",
    "类型": "equipment_type",
    "重量kg": "mass_kg",
    "STA": "sta",
    "WL": "wl",
    "BL": "bl",
    "区域ID": "zone_id",
    "母线ID": "bus_id",
    "功耗kVA": "power_kva_normal",
    "状态": "status",
    "描述": "description",
}

TYPE_MAP = {
    "LRU": "LRU",
    "SRU": "SRU",
    "结构件": "structural",
    "线缆": "cable",
}


async def import_from_excel(
    db: AsyncSession,
    file_path: str,
    zone_map: dict[str, str] | None = None,   # zone_code -> zone_id
    bus_map: dict[str, str] | None = None,     # bus_name -> bus_id
) -> dict:
    """Import equipment from Excel file. Returns summary with success_count, error_rows.

    NOTE: This creates Equipment + ElectricalLoad records only.
    Weight (mass_kg) and config-level attributes (zone, STA/WL/BL, bus assignment)
    should be set via ConfigEquipment when adding equipment to a configuration.
    """
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return {"success_count": 0, "error_rows": [], "total_rows": 0}

    headers = [str(h).strip() if h else "" for h in rows[0]]
    col_idx = {}
    for i, h in enumerate(headers):
        if h in COLUMN_MAP:
            col_idx[COLUMN_MAP[h]] = i

    success_count = 0
    error_rows = []

    for row_num, row in enumerate(rows[1:], start=2):
        try:
            def val(key):
                idx = col_idx.get(key)
                if idx is None or idx >= len(row):
                    return None
                return row[idx]

            part_number = str(val("part_number") or "").strip()
            if not part_number:
                error_rows.append({"row": row_num, "error": "件号为空"})
                continue

            name = str(val("name") or "").strip()
            ata = str(val("ata_chapter") or "").strip()
            etype_raw = str(val("equipment_type") or "LRU").strip()
            etype = TYPE_MAP.get(etype_raw, etype_raw)
            if etype not in ("LRU", "SRU", "structural", "cable"):
                etype = "LRU"
            status_val = str(val("status") or "approved").strip()
            if status_val not in ("in_development", "qualifying", "approved", "discontinued"):
                status_val = "approved"

            equip = Equipment(
                part_number=part_number,
                name=name,
                ata_chapter=ata,
                equipment_type=etype,
                status=status_val,
                description=str(val("description") or "").strip() or None,
            )
            db.add(equip)
            await db.flush()

            # Electrical Load (no bus_id -- that comes from ConfigEquipment.bus_id)
            power = val("power_kva_normal")
            if power is not None:
                el = ElectricalLoad(
                    equipment_id=equip.id,
                    power_kva_normal=float(power),
                )
                db.add(el)

            success_count += 1

        except Exception as exc:
            error_rows.append({"row": row_num, "error": str(exc)})

    await db.commit()
    wb.close()

    return {
        "success_count": success_count,
        "error_rows": error_rows,
        "total_rows": len(rows) - 1,
    }
