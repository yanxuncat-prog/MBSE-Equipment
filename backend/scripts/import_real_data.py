"""Import real CE-25A equipment data from Excel files into AeroEquip platform."""
import asyncio
import re
import sys
import os
import uuid
import random

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import engine, async_session_factory, Base
from app.models import (
    User, Program, Configuration, Equipment, ConfigEquipment,
    WeightBalance, ElectricalLoad, Zone, BusDefinition, Supplier,
    AuditLog,
)
from passlib.context import CryptContext

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

# Try multiple paths: container mount or relative
_candidates = [
    "/app/EXCEL数据",
    os.path.join(os.path.dirname(__file__), "..", "..", "EXCEL数据"),
    os.path.join(os.path.dirname(__file__), "..", "EXCEL数据"),
]
DATA_DIR = next((p for p in _candidates if os.path.isdir(p)), _candidates[0])
print(f"Data directory: {DATA_DIR}")

# Zone mapping: 安装区域文本 -> zone_code + zone_name
ZONE_MAPPING = {
    "机头": ("100", "机头段"),
    "驾驶舱": ("110", "驾驶舱"),
    "下机身前段": ("120", "下机身前段"),
    "右侧前机身": ("130", "右侧前机身"),
    "左侧前机身": ("130", "左侧前机身"),
    "客舱": ("140", "客舱段"),
    "下机身后段": ("150", "下机身后段"),
    "机翼": ("200", "机翼段"),
    "短舱": ("300", "短舱/发动机"),
    "尾段": ("400", "尾段"),
    "设备架": ("110", "驾驶舱"),
}

# ATA chapter cleanup: "EATA23 通信系统" -> "23"
def clean_ata(raw: str) -> str:
    if not raw:
        return "99"
    raw = str(raw).strip()
    m = re.search(r'ATA(\d{2})', raw, re.IGNORECASE)
    if m:
        return m.group(1)
    m = re.match(r'(\d{2})', raw)
    if m:
        return m.group(1)
    return raw[:2] if len(raw) >= 2 else "99"


def get_zone_code(area_text: str) -> tuple[str, str, str]:
    """Returns (zone_code, zone_name, rack_position)."""
    if not area_text:
        return ("999", "未知", "")
    area = str(area_text).strip()
    parts = area.split("-", 1)
    prefix = parts[0]
    detail = parts[1] if len(parts) > 1 else ""

    for key, (code, name) in ZONE_MAPPING.items():
        if key in prefix:
            return (code, name, detail)
    return ("999", "未知", area)


def to_bool(val) -> bool | None:
    """Convert various Chinese/English boolean representations to bool."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ('是', 'Yes', 'yes', 'TRUE', '1'):
        return True
    if s in ('否', 'No', 'no', 'FALSE', '0'):
        return False
    return None


def safe_str(val, max_len: int | None = None) -> str | None:
    """Convert a cell value to a trimmed string, or None if empty/NA."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ('#N/A', 'N/A', 'nan', 'None', '-'):
        return None
    if max_len and len(s) > max_len:
        s = s[:max_len]
    return s


def safe_int(val) -> int | None:
    """Convert a cell value to int, or None."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return int(val)
    s = str(val).strip()
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


# Column index mapping per sheet type for 主清单
# 1&2号构型: no extra col[7] shift
# X号构型: col[7] = "是否使用0号机设备", shifts subsequent columns by 1
COL_MAP = {
    0: {"weight": 11, "area": 5, "is_electric": 13, "dims": 4, "config_cat": 6,
        "layout_adj": 7, "install": 8, "bond_method": 9, "bond_type": 10,
        "has_eicd": 12, "is_primary_elec": 14, "redundancy": 15, "voltage": 16, "power": 17},
    1: {"weight": 12, "area": 5, "is_electric": 14, "dims": 4, "config_cat": 6,
        "use_batch0": 7, "layout_adj": 8, "install": 9, "bond_method": 10, "bond_type": 11,
        "has_eicd": 13, "redundancy": 15, "voltage": 16, "power": 17},
}


async def reset_and_init(db: AsyncSession):
    """Drop all data and recreate."""
    print("Resetting database...")
    async with engine.begin() as conn:
        # Force drop with CASCADE to handle stale FK constraints from old schema
        from sqlalchemy import text
        await conn.execute(text("DROP SCHEMA public CASCADE"))
        await conn.execute(text("CREATE SCHEMA public"))
        await conn.run_sync(Base.metadata.create_all)
    print("Tables recreated.")


async def create_base_data(db: AsyncSession) -> dict:
    """Create users, program, zones, buses. Returns lookup dicts."""
    # Users
    admin = User(username="admin", hashed_password=pwd_context.hash("admin123"),
                 display_name="系统管理员", role="admin")
    engineer = User(username="engineer", hashed_password=pwd_context.hash("eng123"),
                    display_name="设备工程师", role="engineer")
    db.add_all([admin, engineer])
    await db.flush()

    # Program
    program = Program(name="CE-25A", aircraft_type="大型宽体客机",
                      description="CE-25A电动飞机设备管理")
    db.add(program)
    await db.flush()

    # Zones (based on real area data)
    zone_map = {}  # zone_code -> Zone object
    zones_def = [
        ("100", "机头段", 0, 120, 100, 250),
        ("110", "驾驶舱", 120, 350, 150, 280),
        ("120", "下机身前段", 120, 500, 50, 150),
        ("130", "前机身侧面", 350, 500, 100, 280),
        ("140", "客舱段", 350, 800, 150, 300),
        ("150", "下机身后段", 500, 850, 50, 150),
        ("200", "机翼段", 400, 700, 80, 200),
        ("300", "短舱/发动机", 350, 650, 50, 180),
        ("400", "尾段", 850, 1100, 100, 280),
        ("999", "未知", 0, 1100, 0, 300),
    ]
    for code, name, sf, st, wf, wt in zones_def:
        z = Zone(program_id=program.id, zone_code=code, name=name,
                 sta_from=sf, sta_to=st, wl_from=wf, wl_to=wt)
        db.add(z)
        await db.flush()
        zone_map[code] = z

    # Bus definitions (CE-25A electric aircraft)
    bus_map = {}
    buses_def = [
        ("AC BUS 1", "AC", 30.0),
        ("AC BUS 2", "AC", 30.0),
        ("DC BUS 1", "DC", 15.0),
        ("DC BUS 2", "DC", 15.0),
        ("DC ESS", "DC", 8.0),
        ("HOT BAT BUS", "DC", 5.0),
    ]
    for bname, btype, cap in buses_def:
        b = BusDefinition(program_id=program.id, bus_name=bname,
                          bus_type=btype, rated_capacity_kva=cap)
        db.add(b)
        await db.flush()
        bus_map[bname] = b

    await db.commit()
    return {
        "admin": admin, "engineer": engineer,
        "program": program,
        "zone_map": zone_map, "bus_map": bus_map,
    }


# Per-equipment installation context: zone, position, bus assignment
# Stored temporarily during import, then written to ConfigEquipment when creating configs
_equip_install_data: dict[uuid.UUID, dict] = {}  # equipment.id -> {zone_id, sta, wl, bl, rack_position, bus_id, ...}

# Per-equipment config-specific data from 主清单, keyed by (part_number, sheet_idx)
_equip_config_data: dict[tuple[str, int], dict] = {}


async def import_main_list(db: AsyncSession, base: dict) -> dict:
    """Import from 机载系统设备清单_2026.0415.xlsx -- 1&2号构型 and X号构型 sheets."""
    filepath = os.path.join(DATA_DIR, "机载系统设备清单_2026.0415.xlsx")
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return {}

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    zone_map = base["zone_map"]
    equip_by_number = {}  # part_number -> Equipment

    for sheet_idx, sheet_name, config_name in [(0, "1&2号构型", "1/2号机"), (1, "X号构型", "X号机")]:
        ws = wb[wb.sheetnames[sheet_idx]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        cols = COL_MAP[sheet_idx]
        print(f"\nImporting {sheet_name}: {len(rows)} rows (weight col={cols['weight']})")

        success = 0
        skipped = 0
        for row in rows:
            name = str(row[1] or "").strip()
            if not name:
                skipped += 1
                continue

            raw_pn = str(row[2] or "").strip()
            raw_lin = str(row[3] or "").strip()
            part_number = raw_pn or raw_lin or f"UNNAMED-{uuid.uuid4().hex[:8]}"

            # Extract new fields from 主清单
            dims = safe_str(row[cols["dims"]], 100)
            config_cat = safe_str(row[cols["config_cat"]], 20)
            layout_adj = safe_str(row[cols.get("layout_adj", -1)] if cols.get("layout_adj") is not None else None, 500) if "layout_adj" in cols else None
            if "layout_adj" in cols:
                layout_adj = safe_str(row[cols["layout_adj"]], 500)
            install_method = safe_str(row[cols["install"]], 200)
            bond_method = safe_str(row[cols["bond_method"]], 50)
            bond_type = safe_str(row[cols["bond_type"]], 100)
            has_eicd_val = to_bool(row[cols["has_eicd"]])
            is_electric_val = to_bool(row[cols["is_electric"]])
            is_primary_elec_val = to_bool(row[cols.get("is_primary_elec")]) if "is_primary_elec" in cols else None
            redundancy = safe_str(row[cols["redundancy"]], 100)
            voltage = safe_str(row[cols["voltage"]], 50)
            power = safe_str(row[cols["power"]], 50)
            use_batch0 = to_bool(row[cols["use_batch0"]]) if "use_batch0" in cols else None

            # Store config-specific data for later ConfigEquipment creation
            _equip_config_data[(part_number, sheet_idx)] = {
                "install_method": install_method,
                "bonding_method": bond_method,
                "bonding_type": bond_type,
                "layout_adjustment": layout_adj,
                "use_batch0_device": use_batch0,
            }

            # If already imported, try to supplement missing data
            if part_number in equip_by_number:
                existing = equip_by_number[part_number]
                # Supplement LIN号 if missing
                if raw_lin and not existing.lin_number:
                    existing.lin_number = raw_lin
                weight = row[cols["weight"]]
                if weight is not None and isinstance(weight, (int, float)):
                    # Check if existing has no weight -- supplement it
                    from sqlalchemy import select as sel
                    wb_check = await db.execute(sel(WeightBalance).where(WeightBalance.equipment_id == existing.id))
                    if wb_check.scalar_one_or_none() is None:
                        try:
                            wb_new = WeightBalance(equipment_id=existing.id, mass_kg=float(weight))
                            db.add(wb_new)
                            await db.flush()
                        except Exception:
                            pass
                # Supplement equipment-level fields from second sheet if not set
                if dims and not existing.dimensions_mm:
                    existing.dimensions_mm = dims
                if config_cat and not existing.config_category:
                    existing.config_category = config_cat
                if has_eicd_val is not None and existing.has_eicd is None:
                    existing.has_eicd = has_eicd_val
                if is_electric_val is not None and existing.is_electrical is None:
                    existing.is_electrical = is_electric_val
                if is_primary_elec_val is not None and existing.is_primary_electrical is None:
                    existing.is_primary_electrical = is_primary_elec_val
                if redundancy and not existing.power_redundancy:
                    existing.power_redundancy = redundancy
                if voltage and not existing.power_voltage:
                    existing.power_voltage = voltage
                if power and not existing.power_watts:
                    existing.power_watts = power
                success += 1
                continue

            ata = clean_ata(str(row[0] or ""))
            area = str(row[cols["area"]] or "").strip()
            zone_code, zone_name, rack_pos = get_zone_code(area)
            weight = row[cols["weight"]]
            is_electric = str(row[cols["is_electric"]] or "").strip()

            equip = Equipment(
                part_number=part_number,
                name=name,
                lin_number=raw_lin or None,
                ata_chapter=ata,
                equipment_type="LRU" if is_electric == "是" else "structural",
                status="approved",
                description=f"区域: {area}" if area else None,
                # New fields from 主清单
                dimensions_mm=dims,
                config_category=config_cat,
                has_eicd=has_eicd_val,
                is_electrical=is_electric_val,
                is_primary_electrical=is_primary_elec_val,
                power_redundancy=redundancy,
                power_voltage=voltage,
                power_watts=power,
            )
            db.add(equip)
            try:
                await db.flush()
            except Exception as e:
                await db.rollback()
                skipped += 1
                continue

            # Compute installation position data (will be stored on ConfigEquipment later)
            zone_obj = zone_map.get(zone_code, zone_map.get("999"))
            approx_sta = (zone_obj.sta_from + zone_obj.sta_to) / 2 if zone_obj else 500
            approx_wl = (zone_obj.wl_from + zone_obj.wl_to) / 2 if zone_obj else 180
            sta_offset = random.uniform(-50, 50)
            wl_offset = random.uniform(-20, 20)

            _equip_install_data[equip.id] = {
                "zone_id": zone_obj.id if zone_obj else None,
                "sta": approx_sta + sta_offset,
                "wl": approx_wl + wl_offset,
                "bl": random.uniform(-30, 30),
                "rack_position": rack_pos if rack_pos else None,
            }

            # Weight (skip #N/A, strings, etc.)
            if weight is not None and isinstance(weight, (int, float)):
                try:
                    mass = float(weight)
                    if mass > 0:
                        wb_obj = WeightBalance(
                            equipment_id=equip.id,
                            mass_kg=mass,
                        )
                        db.add(wb_obj)
                except (ValueError, TypeError):
                    pass

            equip_by_number[part_number] = equip
            success += 1

        await db.commit()
        print(f"  Imported: {success}, Skipped: {skipped}")

    wb.close()
    return equip_by_number


async def import_supplier_and_ata_data(db: AsyncSession, equip_map: dict, base: dict):
    """Supplement equipment with ALL fields from 0号机 ATA sheet."""
    filepath = os.path.join(DATA_DIR, "CE-25A 0号机_2026-04-21.xlsx")
    if not os.path.exists(filepath):
        print("0号机 file not found, skipping supplier/ATA import")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]  # ATA章节设备表
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"\nImporting supplier + ATA fields from 0号机: {len(rows)} rows")

    # Collect unique suppliers
    supplier_cache = {}  # name -> Supplier object
    updated = 0

    for row in rows:
        doors_number = str(row[1] or "").strip()  # Col[1] DOORS编号 = our part_number
        supplier_name = str(row[7] or "").strip()

        if not doors_number:
            continue

        # Find matching equipment by DOORS number
        equip = equip_map.get(doors_number)
        if not equip:
            continue

        # Supplier
        if supplier_name and supplier_name != "TBD":
            if supplier_name not in supplier_cache:
                s = Supplier(name=supplier_name)
                db.add(s)
                await db.flush()
                supplier_cache[supplier_name] = s
            equip.supplier_id = supplier_cache[supplier_name].id

        # Col[0] 内部设备编号
        val = safe_str(row[0], 50)
        if val and not equip.internal_number:
            equip.internal_number = val

        # Col[2] LIN号
        val = safe_str(row[2], 50)
        if val and not equip.lin_number:
            equip.lin_number = val

        # Col[4] 英文名称
        val = safe_str(row[4], 300)
        if val and not equip.name_en:
            equip.name_en = val

        # Col[5] 英文缩写
        val = safe_str(row[5], 50)
        if val and not equip.abbreviation_en:
            equip.abbreviation_en = val

        # Col[6] 供应商件号
        val = safe_str(row[6], 100)
        if val and not equip.supplier_part_number:
            equip.supplier_part_number = val

        # Col[10] DAL
        val = safe_str(row[10], 5)
        if val and not equip.dal:
            equip.dal = val

        # Col[11] 壳体是否金属
        bval = to_bool(row[11])
        if bval is not None and equip.is_metal_shell is None:
            equip.is_metal_shell = bval

        # Col[12] 金属壳体是否经特殊处理不易导电
        val = safe_str(row[12], 200)
        if val and not equip.metal_shell_non_conductive:
            equip.metal_shell_non_conductive = val

        # Col[13] 设备内共地情况
        val = safe_str(row[13], 200)
        if val and not equip.internal_grounding:
            equip.internal_grounding = val

        # Col[14] 壳体接地方式
        val = safe_str(row[14], 100)
        if val and not equip.shell_grounding_method:
            equip.shell_grounding_method = val

        # Col[15] 壳体接地是否故障电流路径
        val = safe_str(row[15], 200)
        if val and not equip.shell_grounding_fault_path:
            equip.shell_grounding_fault_path = val

        # Col[16] 其他接地特殊要求
        val = safe_str(row[16])
        if val and not equip.grounding_special_requirements:
            equip.grounding_special_requirements = val

        # Col[17] 连接器或接线柱数量
        ival = safe_int(row[17])
        if ival is not None and equip.connector_count is None:
            equip.connector_count = ival

        # Col[18] 是否选装设备
        bval = to_bool(row[18])
        if bval is not None and equip.is_optional is None:
            equip.is_optional = bval

        # Col[19] 是否有特殊布线需求
        bval = to_bool(row[19])
        if bval is not None and equip.has_special_wiring is None:
            equip.has_special_wiring = bval

        # Col[20] 装机架次
        val = safe_str(row[20], 50)
        if val and not equip.aircraft_batch:
            equip.aircraft_batch = val

        # Col[21] 设备负责人
        val = safe_str(row[21], 50)
        if val and not equip.responsible_person:
            equip.responsible_person = val

        # Col[22] 设备等级
        val = safe_str(row[22], 50)
        if val and not equip.equipment_level:
            equip.equipment_level = val

        # Col[23] 正常工作电压范围(V)
        val = safe_str(row[23], 100)
        if val and not equip.voltage_range:
            equip.voltage_range = val

        # Col[24] 设备物理特性
        val = safe_str(row[24])
        if val and not equip.physical_characteristics:
            equip.physical_characteristics = val

        # Col[25] 备注
        val = safe_str(row[25])
        if val and not equip.notes:
            equip.notes = val

        updated += 1

    await db.commit()
    print(f"  Suppliers created: {len(supplier_cache)}, Equipment updated: {updated}")
    wb.close()


async def import_temperature_data(db: AsyncSession, equip_map: dict):
    """Supplement equipment with DO-160 temperature qualification data."""
    filepath = os.path.join(DATA_DIR, "CE-25A飞机0号机有人首飞机载系统设备清单_0417—温度-汇总V1.xlsx")
    if not os.path.exists(filepath):
        print("温度汇总 file not found, skipping temperature import")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f"\nImporting temperature data: {len(rows)} rows")

    # Determine column layout by reading header
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    # Build a col index map from header names
    col_idx = {}
    for i, h in enumerate(header_row):
        if h:
            col_idx[str(h).strip()] = i

    # Try to find the DOORS编号 column (col[2] per spec)
    doors_col = col_idx.get("设备编号（DOORS）", col_idx.get("设备编号(DOORS)", 2))

    updated = 0
    for row in rows:
        if len(row) < 5:
            continue
        doors_number = str(row[doors_col] or "").strip() if doors_col < len(row) else ""
        if not doors_number:
            continue

        equip = equip_map.get(doors_number)
        if not equip:
            continue

        # Map columns by header name or fall back to positional
        def get_col(name, fallback=None):
            idx = col_idx.get(name, fallback)
            if idx is not None and idx < len(row):
                return row[idx]
            return None

        # 首飞是否上机
        bval = to_bool(get_col("首飞是否上机"))
        if bval is not None and equip.first_flight_onboard is None:
            equip.first_flight_onboard = bval

        # 二阶段是否上机
        bval = to_bool(get_col("二阶段是否上机"))
        if bval is not None and equip.phase2_onboard is None:
            equip.phase2_onboard = bval

        # 设计要求等级
        val = safe_str(get_col("设计要求等级"), 100)
        if val and not equip.do160_temp_design_level:
            equip.do160_temp_design_level = val

        # DO-160第4章温度高度鉴定等级
        val = safe_str(get_col("DO-160第4章温度高度鉴定等级", col_idx.get("DO-160第4章温度鉴定等级")), 100)
        if val and not equip.do160_temp_qual_level:
            equip.do160_temp_qual_level = val

        # 鉴定工作温度范围
        val = safe_str(get_col("鉴定工作温度范围"), 100)
        if val and not equip.do160_temp_qual_range:
            equip.do160_temp_qual_range = val

        # 鉴定符合情况
        val = safe_str(get_col("鉴定符合情况"), 200)
        if val and not equip.do160_temp_compliance:
            equip.do160_temp_compliance = val

        # 正常工作温度(℃)
        val = safe_str(get_col("正常工作温度(℃)", col_idx.get("正常工作温度")), 50)
        if val and not equip.normal_operating_temp:
            equip.normal_operating_temp = val

        # 短时工作温度(℃)
        val = safe_str(get_col("短时工作温度(℃)", col_idx.get("短时工作温度")), 50)
        if val and not equip.short_term_temp:
            equip.short_term_temp = val

        # 地面停放温度(℃)
        val = safe_str(get_col("地面停放温度(℃)", col_idx.get("地面停放温度")), 50)
        if val and not equip.ground_storage_temp:
            equip.ground_storage_temp = val

        # 高度(m)
        val = safe_str(get_col("高度(m)", col_idx.get("高度")), 50)
        if val and not equip.operating_altitude:
            equip.operating_altitude = val

        # 鉴定报告编号
        val = safe_str(get_col("鉴定报告编号"), 200)
        if val and not equip.qual_report_number:
            equip.qual_report_number = val

        updated += 1

    await db.commit()
    print(f"  Temperature data updated: {updated} equipment")
    wb.close()


async def import_bonding_data(db: AsyncSession, equip_map: dict):
    """Import 电搭接 bonding data from the 0号机 file's 电搭接情况统计 sheet."""
    filepath = os.path.join(DATA_DIR, "CE-25A 0号机_2026-04-21.xlsx")
    if not os.path.exists(filepath):
        print("0号机 file not found, skipping bonding import")
        return

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

    # Look for 电搭接情况统计 sheet
    bonding_sheet = None
    for sn in wb.sheetnames:
        if "电搭接" in sn or "搭接" in sn:
            bonding_sheet = sn
            break

    if not bonding_sheet:
        print("  No 电搭接情况统计 sheet found, skipping")
        wb.close()
        return

    ws = wb[bonding_sheet]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        wb.close()
        return

    # Read header to find columns
    header = rows[0]
    col_idx = {}
    for i, h in enumerate(header):
        if h:
            col_idx[str(h).strip()] = i

    data_rows = rows[1:]
    print(f"\nImporting bonding data from '{bonding_sheet}': {len(data_rows)} rows")

    # Find the DOORS编号 column
    doors_col = None
    for key in ["设备编号（DOORS）", "设备编号(DOORS)", "DOORS编号"]:
        if key in col_idx:
            doors_col = col_idx[key]
            break
    if doors_col is None:
        # Fall back to col index 2
        doors_col = 2

    # Find bonding-specific columns
    resistance_col = None
    position_col = None
    pace_col = None
    install_diagram_col = None

    for key in col_idx:
        if "阻值" in key and "mΩ" in key:
            resistance_col = col_idx[key]
        elif "阻值" in key:
            resistance_col = col_idx[key]
        if "搭接位置" in key or "结构零件号" in key:
            position_col = col_idx[key]
        if "PACE" in key:
            pace_col = col_idx[key]
        if "安装图" in key:
            install_diagram_col = col_idx[key]

    # Store bonding data keyed by part_number for later ConfigEquipment update
    global _bonding_data
    _bonding_data = {}  # part_number -> {bonding_resistance, bonding_position, in_pace_drawing}

    updated = 0
    for row in data_rows:
        if len(row) <= doors_col:
            continue
        doors_number = str(row[doors_col] or "").strip()
        if not doors_number:
            continue

        equip = equip_map.get(doors_number)
        if not equip:
            continue

        bonding_info = {}
        if resistance_col is not None and resistance_col < len(row):
            bonding_info["bonding_resistance"] = safe_str(row[resistance_col], 50)
        if position_col is not None and position_col < len(row):
            bonding_info["bonding_position"] = safe_str(row[position_col], 200)
        if pace_col is not None and pace_col < len(row):
            bonding_info["in_pace_drawing"] = to_bool(row[pace_col])

        if any(v is not None for v in bonding_info.values()):
            _bonding_data[doors_number] = bonding_info
            updated += 1

    await db.commit()
    print(f"  Bonding data collected: {updated} equipment")
    wb.close()


# Global bonding data store
_bonding_data: dict[str, dict] = {}


async def create_configurations(db: AsyncSession, equip_map: dict, base: dict):
    """Create configurations and assign equipment via ConfigEquipment with position/bus data."""
    program = base["program"]
    admin = base["admin"]
    bus_map = base["bus_map"]

    # Load all equipment
    result = await db.execute(select(Equipment))
    all_equip = result.scalars().all()
    print(f"\nTotal equipment in DB: {len(all_equip)}")

    # Build reverse map: equipment.id -> part_number
    id_to_pn = {e.id: e.part_number for e in all_equip}

    # Assign bus IDs round-robin for electric equipment
    bus_list = list(bus_map.values())

    configs_to_create = [
        ("V1.0-基线", "baseline", "1/2号机基线构型", 0),
        ("V1.1-0号机", "draft", "0号机构型（首飞）", 0),
        ("V2.0-X号机", "draft", "X号机构型", 1),
    ]

    for version, status, desc, sheet_idx in configs_to_create:
        config = Configuration(
            program_id=program.id,
            version=version,
            status=status,
            description=desc,
            created_by=admin.id,
        )
        db.add(config)
        await db.flush()

        for idx, equip in enumerate(all_equip):
            install = _equip_install_data.get(equip.id, {})
            pn = id_to_pn.get(equip.id, "")

            # Assign a bus for electric LRU equipment (round-robin)
            bus_id = None
            if equip.equipment_type == "LRU" and bus_list:
                bus_id = bus_list[idx % len(bus_list)].id

            # Get config-specific data from 主清单
            cfg_data = _equip_config_data.get((pn, sheet_idx), {})
            # Also try the other sheet if not found
            if not cfg_data:
                cfg_data = _equip_config_data.get((pn, 1 - sheet_idx), {})

            # Get bonding data
            bonding = _bonding_data.get(pn, {})

            ce = ConfigEquipment(
                config_id=config.id,
                equipment_id=equip.id,
                zone_id=install.get("zone_id"),
                sta=install.get("sta"),
                wl=install.get("wl"),
                bl=install.get("bl"),
                rack_position=install.get("rack_position"),
                bus_id=bus_id,
                # New config-specific fields
                install_method=cfg_data.get("install_method"),
                bonding_method=cfg_data.get("bonding_method"),
                bonding_type=cfg_data.get("bonding_type"),
                layout_adjustment=cfg_data.get("layout_adjustment"),
                use_batch0_device=cfg_data.get("use_batch0_device"),
                bonding_resistance=bonding.get("bonding_resistance"),
                bonding_position=bonding.get("bonding_position"),
                in_pace_drawing=bonding.get("in_pace_drawing"),
            )
            db.add(ce)

        await db.flush()
        print(f"  Config '{version}' ({status}): {len(all_equip)} equipment")

    await db.commit()


async def main():
    print("=" * 60)
    print("AeroEquip -- Real CE-25A Data Import (Full Fields)")
    print("=" * 60)

    await reset_and_init(None)

    async with async_session_factory() as db:
        base = await create_base_data(db)
        equip_map = await import_main_list(db, base)
        await import_supplier_and_ata_data(db, equip_map, base)
        await import_temperature_data(db, equip_map)
        await import_bonding_data(db, equip_map)
        await create_configurations(db, equip_map, base)

    # Final stats
    async with async_session_factory() as db:
        equip_count = (await db.execute(select(Equipment))).scalars().all()
        wb_count = (await db.execute(select(WeightBalance))).scalars().all()
        ce_count = (await db.execute(select(ConfigEquipment))).scalars().all()
        sup_count = (await db.execute(select(Supplier))).scalars().all()
        config_count = (await db.execute(select(Configuration))).scalars().all()
        zone_count = (await db.execute(select(Zone))).scalars().all()

        # Count how many equipment have temperature data
        temp_count = sum(1 for e in equip_count if e.do160_temp_design_level or e.normal_operating_temp)
        ata_count = sum(1 for e in equip_count if e.name_en or e.dal)
        bonding_count = sum(1 for c in ce_count if c.bonding_resistance or c.bonding_method)

    print("\n" + "=" * 60)
    print("Import Complete!")
    print(f"  设备: {len(equip_count)}")
    print(f"  构型设备关联: {len(ce_count)}")
    print(f"  重量数据: {len(wb_count)}")
    print(f"  供应商: {len(sup_count)}")
    print(f"  区域: {len(zone_count)}")
    print(f"  构型: {len(config_count)}")
    print(f"  含ATA详细字段(英文名/DAL): {ata_count}")
    print(f"  含温度鉴定数据: {temp_count}")
    print(f"  含电搭接数据: {bonding_count}")
    print("=" * 60)


if __name__ == "__main__":
    asyncio.run(main())
